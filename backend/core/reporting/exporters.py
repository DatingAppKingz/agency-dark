"""
Report export handlers for different formats.
"""
from typing import Dict, List, Optional, Any
from datetime import datetime
import pandas as pd
from io import BytesIO
import json
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
import base64
import logging

logger = logging.getLogger(__name__)


class PDFExporter:
    """Export reports to PDF format."""
    
    def __init__(self):
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()
    
    def _setup_custom_styles(self):
        """Setup custom PDF styles."""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#2E86AB'),
            spaceAfter=30
        ))
        
        self.styles.add(ParagraphStyle(
            name='CustomHeading',
            parent=self.styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#2E86AB'),
            spaceAfter=12
        ))
    
    async def export(
        self,
        data: pd.DataFrame,
        report_name: str,
        metadata: Dict[str, Any],
        charts: Optional[List[Dict[str, Any]]] = None
    ) -> BytesIO:
        """Export data to PDF."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        story = []
        
        # Add title
        story.append(Paragraph(report_name, self.styles['CustomTitle']))
        story.append(Spacer(1, 12))
        
        # Add metadata
        if metadata:
            story.append(Paragraph("Report Information", self.styles['CustomHeading']))
            meta_data = [
                ["Generated:", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")],
                ["Records:", f"{len(data):,}"],
                ["Period:", metadata.get("period", "All time")]
            ]
            meta_table = Table(meta_data, colWidths=[2*inch, 4*inch])
            meta_table.setStyle(TableStyle([
                ('FONT', (0, 0), (-1, -1), 'Helvetica', 10),
                ('TEXTCOLOR', (0, 0), (0, -1), colors.grey),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            story.append(meta_table)
            story.append(Spacer(1, 20))
        
        # Add charts if provided
        if charts:
            for chart_config in charts:
                chart_img = await self._create_chart(data, chart_config)
                if chart_img:
                    story.append(Image(chart_img, width=6*inch, height=4*inch))
                    story.append(Spacer(1, 20))
        
        # Add data table
        story.append(Paragraph("Data", self.styles['CustomHeading']))
        
        # Convert DataFrame to table data
        table_data = [data.columns.tolist()]  # Headers
        for _, row in data.iterrows():
            table_data.append([str(val) for val in row.values])
        
        # Limit rows for PDF
        if len(table_data) > 101:  # 100 data rows + header
            table_data = table_data[:101]
            story.append(Paragraph(
                f"<i>Showing first 100 of {len(data):,} records</i>",
                self.styles['Italic']
            ))
        
        # Create table
        col_widths = [len(str(col)) * 7 for col in data.columns]
        max_width = 7 * inch
        total_width = sum(col_widths)
        if total_width > max_width:
            scale = max_width / total_width
            col_widths = [w * scale for w in col_widths]
        
        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            # Header style
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E86AB')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            
            # Data style
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
        ]))
        
        story.append(table)
        
        # Build PDF
        doc.build(story)
        buffer.seek(0)
        return buffer
    
    async def _create_chart(
        self,
        data: pd.DataFrame,
        chart_config: Dict[str, Any]
    ) -> Optional[BytesIO]:
        """Create chart image for PDF."""
        try:
            plt.figure(figsize=(10, 6))
            
            chart_type = chart_config.get("type", "bar")
            x_column = chart_config.get("x_column")
            y_column = chart_config.get("y_column")
            
            if not x_column or not y_column:
                return None
            
            if chart_type == "bar":
                plt.bar(data[x_column], data[y_column])
            elif chart_type == "line":
                plt.plot(data[x_column], data[y_column])
            elif chart_type == "pie":
                plt.pie(data[y_column], labels=data[x_column], autopct='%1.1f%%')
            
            plt.title(chart_config.get("title", ""))
            plt.xlabel(x_column)
            plt.ylabel(y_column)
            plt.xticks(rotation=45)
            plt.tight_layout()
            
            # Save to buffer
            buffer = BytesIO()
            plt.savefig(buffer, format='png')
            buffer.seek(0)
            plt.close()
            
            return buffer
            
        except Exception as e:
            logger.error(f"Error creating chart: {e}")
            return None


class ExcelExporter:
    """Export reports to Excel format with formatting."""
    
    def __init__(self):
        self.header_font = Font(bold=True, color="FFFFFF", size=12)
        self.header_fill = PatternFill(start_color="2E86AB", end_color="2E86AB", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center")
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    async def export(
        self,
        data: pd.DataFrame,
        report_name: str,
        metadata: Dict[str, Any],
        charts: Optional[List[Dict[str, Any]]] = None
    ) -> BytesIO:
        """Export data to Excel with formatting."""
        buffer = BytesIO()
        
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            # Write main data
            data.to_excel(writer, sheet_name='Data', index=False, startrow=3)
            
            # Get workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Data']
            
            # Add report title
            worksheet['A1'] = report_name
            worksheet['A1'].font = Font(bold=True, size=16, color="2E86AB")
            worksheet.merge_cells(f'A1:{chr(65 + len(data.columns) - 1)}1')
            
            # Add metadata
            worksheet['A2'] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
            worksheet['A2'].font = Font(italic=True, size=10)
            
            # Format headers
            for col_num, column in enumerate(data.columns, 1):
                cell = worksheet.cell(row=4, column=col_num)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = self.header_alignment
                cell.border = self.border
                
                # Auto-adjust column width
                max_length = max(
                    len(str(column)),
                    data[column].astype(str).map(len).max()
                )
                worksheet.column_dimensions[chr(64 + col_num)].width = min(max_length + 2, 50)
            
            # Format data cells
            for row in range(5, len(data) + 5):
                for col in range(1, len(data.columns) + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.border = self.border
                    
                    # Alternate row colors
                    if row % 2 == 0:
                        cell.fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
            
            # Add summary sheet
            if metadata.get("include_summary", True):
                summary_sheet = workbook.create_sheet("Summary")
                self._add_summary_sheet(summary_sheet, data, metadata)
            
            # Add charts if requested
            if charts:
                chart_sheet = workbook.create_sheet("Charts")
                await self._add_charts(chart_sheet, data, charts)
        
        buffer.seek(0)
        return buffer
    
    def _add_summary_sheet(
        self,
        worksheet,
        data: pd.DataFrame,
        metadata: Dict[str, Any]
    ):
        """Add summary statistics sheet."""
        worksheet['A1'] = "Report Summary"
        worksheet['A1'].font = Font(bold=True, size=14)
        
        row = 3
        for column in data.select_dtypes(include=['number']).columns:
            worksheet[f'A{row}'] = column
            worksheet[f'B{row}'] = "Count"
            worksheet[f'C{row}'] = len(data[column])
            worksheet[f'B{row+1}'] = "Mean"
            worksheet[f'C{row+1}'] = data[column].mean()
            worksheet[f'B{row+2}'] = "Min"
            worksheet[f'C{row+2}'] = data[column].min()
            worksheet[f'B{row+3}'] = "Max"
            worksheet[f'C{row+3}'] = data[column].max()
            row += 5
    
    async def _add_charts(
        self,
        worksheet,
        data: pd.DataFrame,
        charts: List[Dict[str, Any]]
    ):
        """Add charts to Excel."""
        # TODO: Implement Excel chart generation
        pass


class CSVExporter:
    """Export reports to CSV format."""
    
    async def export(
        self,
        data: pd.DataFrame,
        report_name: str,
        metadata: Dict[str, Any]
    ) -> BytesIO:
        """Export data to CSV."""
        buffer = BytesIO()
        
        # Add metadata as comments if requested
        if metadata.get("include_metadata", False):
            buffer.write(f"# Report: {report_name}\n".encode())
            buffer.write(f"# Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}\n".encode())
            buffer.write(f"# Records: {len(data)}\n".encode())
            buffer.write("#\n".encode())
        
        # Write data
        data.to_csv(buffer, index=False, encoding='utf-8')
        buffer.seek(0)
        return buffer


class HTMLExporter:
    """Export reports to HTML format."""
    
    def __init__(self):
        self.template = """
        <!DOCTYPE html>
        <html>
        <head>
            <title>{title}</title>
            <meta charset="utf-8">
            <style>
                body {{ font-family: Arial, sans-serif; margin: 20px; }}
                h1 {{ color: #2E86AB; }}
                .metadata {{ background: #f5f5f5; padding: 10px; margin-bottom: 20px; }}
                table {{ border-collapse: collapse; width: 100%; }}
                th {{ background: #2E86AB; color: white; padding: 10px; text-align: left; }}
                td {{ border: 1px solid #ddd; padding: 8px; }}
                tr:nth-child(even) {{ background: #f5f5f5; }}
                .chart {{ margin: 20px 0; }}
            </style>
        </head>
        <body>
            <h1>{title}</h1>
            <div class="metadata">
                <p>Generated: {generated_at}</p>
                <p>Records: {record_count:,}</p>
            </div>
            {charts}
            <h2>Data</h2>
            {table}
        </body>
        </html>
        """
    
    async def export(
        self,
        data: pd.DataFrame,
        report_name: str,
        metadata: Dict[str, Any],
        charts: Optional[List[Dict[str, Any]]] = None
    ) -> BytesIO:
        """Export data to HTML."""
        # Generate table HTML
        table_html = data.to_html(
            index=False,
            classes=['report-table'],
            table_id='data-table',
            escape=False
        )
        
        # Generate charts HTML
        charts_html = ""
        if charts:
            for chart_config in charts:
                chart_data = await self._create_chart_html(data, chart_config)
                if chart_data:
                    charts_html += f'<div class="chart">{chart_data}</div>'
        
        # Generate final HTML
        html = self.template.format(
            title=report_name,
            generated_at=datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC'),
            record_count=len(data),
            charts=charts_html,
            table=table_html
        )
        
        buffer = BytesIO(html.encode('utf-8'))
        return buffer
    
    async def _create_chart_html(
        self,
        data: pd.DataFrame,
        chart_config: Dict[str, Any]
    ) -> Optional[str]:
        """Create chart HTML using Chart.js or similar."""
        # TODO: Implement interactive charts
        return None


class JSONExporter:
    """Export reports to JSON format."""
    
    async def export(
        self,
        data: pd.DataFrame,
        report_name: str,
        metadata: Dict[str, Any]
    ) -> BytesIO:
        """Export data to JSON."""
        result = {
            "report": {
                "name": report_name,
                "generated_at": datetime.utcnow().isoformat(),
                "metadata": metadata,
                "record_count": len(data)
            },
            "data": data.to_dict(orient='records')
        }
        
        json_str = json.dumps(result, indent=2, default=str)
        buffer = BytesIO(json_str.encode('utf-8'))
        return buffer


class ReportExportManager:
    """Manager for handling all export formats."""
    
    def __init__(self):
        self.exporters = {
            "pdf": PDFExporter(),
            "excel": ExcelExporter(),
            "csv": CSVExporter(),
            "html": HTMLExporter(),
            "json": JSONExporter()
        }
    
    async def export(
        self,
        data: pd.DataFrame,
        format: str,
        report_name: str,
        metadata: Optional[Dict[str, Any]] = None,
        charts: Optional[List[Dict[str, Any]]] = None
    ) -> BytesIO:
        """Export report in specified format."""
        exporter = self.exporters.get(format.lower())
        if not exporter:
            raise ValueError(f"Unsupported export format: {format}")
        
        metadata = metadata or {}
        
        if format in ["pdf", "excel", "html"] and charts:
            return await exporter.export(data, report_name, metadata, charts)
        else:
            return await exporter.export(data, report_name, metadata)


# Singleton instance
export_manager = ReportExportManager()